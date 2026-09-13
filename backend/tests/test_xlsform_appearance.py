"""`appearance: field-list` on a group survives the import.

Both engines and `build_screen_plan` have implemented §11.1's field-list screen
since the partition existed. The importer did not read the column, so **no
workbook could ever ask for one**: every imported form was one question per
screen, always, and nothing said so — `appearance` was on the known-ignored list
with the reason "the IR carries `appearance` but no client reads it yet", which
was true of the column and not of the value.

What made it indefensible is a number. The Sindh-scale run imported 2,128
questions to 2,101 screens, so an enumerator taps 2,101 times through a
questionnaire whose paper original groups its questions into blocks
(`docs/scale-run-2026-09-12-sindh.md` §4.1). Telling RCons to emit a column we
discard would have been asking for a workbook we could not honour.

The last two tests are the pair that matters: a repeat inside a field-list group
is a §10.2 contradiction, and now that a workbook can *produce* one, the
importer has to report it against a row rather than let it surface as a compile
failure with no provenance.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl

from app.modules.form_engine.screens import build_screen_plan
from app.modules.forms.xlsform.importer import import_workbook

COLUMNS = ["type", "name", "label::English (en)", "appearance"]


def _workbook(rows: list[list[str]]) -> bytes:
    book = openpyxl.Workbook()
    survey = book.active
    survey.title = "survey"
    survey.append(COLUMNS)
    for row in rows:
        survey.append(row + [""] * (len(COLUMNS) - len(row)))
    choices = book.create_sheet("choices")
    choices.append(["list_name", "name", "label::English (en)"])
    choices.append(["yes_no", "yes", "Yes"])
    choices.append(["yes_no", "no", "No"])
    settings = book.create_sheet("settings")
    settings.append(["form_title", "form_id", "version", "default_language"])
    settings.append(["Appearance", "appearance_test", "1", "English (en)"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _node(form: dict[str, Any], node_id: str) -> dict[str, Any]:
    def walk(nodes: list[dict[str, Any]]) -> dict[str, Any] | None:
        for node in nodes:
            if node.get("id") == node_id:
                return node
            found = walk(node.get("children", []))
            if found is not None:
                return found
        return None

    found = walk(form.get("children", []))
    assert found is not None, node_id
    return found


def test_a_field_list_group_becomes_one_screen() -> None:
    result = import_workbook(
        _workbook(
            [
                ["begin group", "block", "Block", "field-list"],
                ["text", "q1", "One"],
                ["integer", "q2", "Two"],
                ["select_one yes_no", "q3", "Three"],
                ["end group"],
                ["text", "after", "After"],
            ]
        )
    )
    assert [d.message for d in result.diagnostics if d.severity == "error"] == []
    assert _node(result.form, "block")["appearance"] == "field-list"

    plan = build_screen_plan(result.form)
    assert len(plan.screens) == 2, "three questions on one screen, then `after`"
    assert plan.screens[0].question_ids == ("q1", "q2", "q3")
    assert plan.screens[0].group_id == "block"


def test_without_the_column_it_is_still_one_screen_per_question() -> None:
    """The control. Without it the test above is satisfied by a planner that
    puts every group on one screen, which is a different and wrong platform."""
    result = import_workbook(
        _workbook(
            [
                ["begin group", "block", "Block"],
                ["text", "q1", "One"],
                ["integer", "q2", "Two"],
                ["select_one yes_no", "q3", "Three"],
                ["end group"],
            ]
        )
    )
    assert "appearance" not in _node(result.form, "block")
    assert len(build_screen_plan(result.form).screens) == 3


def test_extra_tokens_beside_field_list_are_kept_and_reported() -> None:
    """Real workbooks write `field-list minimal`, so the cell is read as tokens.

    Comparing the whole cell would drop the grouping because of a display hint
    sitting next to it — which is the silent half of the failure this test
    exists for.
    """
    result = import_workbook(
        _workbook(
            [
                ["begin group", "block", "Block", "field-list minimal"],
                ["text", "q1", "One"],
                ["text", "q2", "Two"],
                ["end group"],
            ]
        )
    )
    assert _node(result.form, "block")["appearance"] == "field-list"
    assert len(build_screen_plan(result.form).screens) == 1
    notes = [d.message for d in result.diagnostics if d.code == "appearance_not_read"]
    assert len(notes) == 1 and "minimal" in notes[0]


def test_table_list_groups_and_says_which_half_was_taken() -> None:
    result = import_workbook(
        _workbook(
            [
                ["begin group", "block", "Block", "table-list"],
                ["select_one yes_no", "q1", "One"],
                ["select_one yes_no", "q2", "Two"],
                ["end group"],
            ]
        )
    )
    assert _node(result.form, "block")["appearance"] == "field-list"
    assert len(build_screen_plan(result.form).screens) == 1
    assert any(d.code == "appearance_table_list" for d in result.diagnostics)


def test_field_list_on_a_repeat_is_not_applied_and_says_so() -> None:
    """§11.3 already makes a repeat one screen entered and left.

    Applying it would assert the contradiction §10.2 refuses instead of
    reporting it.
    """
    result = import_workbook(
        _workbook(
            [
                ["integer", "n", "How many"],
                ["begin repeat", "members", "Members", "field-list"],
                ["text", "member_name", "Name"],
                ["end repeat"],
            ]
        )
    )
    assert "appearance" not in _node(result.form, "members")
    assert any(d.code == "appearance_on_repeat" for d in result.diagnostics)
    assert result.publishable


def test_a_repeat_inside_a_field_list_group_is_now_reachable_and_refused() -> None:
    """The §10.2 contradiction, which no workbook could express until now.

    `test_repeat_in_field_list.py` has held both engines to this refusal since
    it was written, against IR built by hand. This is the first time an *import*
    can produce the shape, so it is the first time the importer has to report it
    against a row instead of letting it reach the publish endpoint.
    """
    result = import_workbook(
        _workbook(
            [
                ["begin group", "block", "Block", "field-list"],
                ["text", "q1", "One"],
                ["begin repeat", "members", "Members"],
                ["text", "member_name", "Name"],
                ["end repeat"],
                ["end group"],
            ]
        )
    )
    assert not result.publishable
    refusals = [d for d in result.diagnostics if d.severity == "error"]
    assert refusals, "the contradiction must be reported, not published"
    assert any("field-list" in d.message for d in refusals)
    assert any(d.ref is not None for d in refusals), (
        "an author with the file open should not have to hunt for the row"
    )
