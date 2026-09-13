"""Runs every answerability vector against the Python reference.

The Kotlin engine runs the same files from shared/form-engine. A form that
publishes on one implementation and is refused on the other is a release
blocker: a form author would meet a refusal their builder told them was not
there. Spec: Form IR §10.2, §2.1, §6.2.

The second test is the one that matters as much as the first, for the reason
`conformance/reachability`'s runner gives: a rule both engines implement and no
publish path calls is a rule that refuses nothing. Defect 31 is what a rule
called from nowhere looks like — the note published, deployed, and reached a
phone.
"""

from __future__ import annotations

import json
import pathlib

import pytest

from app.modules.form_engine.answerability import check_answerability
from app.modules.forms.service import PublishRefused, check_publishable

VECTOR_DIR = pathlib.Path(__file__).resolve().parents[2] / "conformance" / "answerability"
VECTORS = sorted(VECTOR_DIR.glob("*.json"))

assert VECTORS, f"no answerability vectors found in {VECTOR_DIR}"


@pytest.mark.parametrize("path", VECTORS, ids=lambda p: p.stem)
def test_vector(path: pathlib.Path) -> None:
    vector = json.loads(path.read_text())
    assert vector["type"] == "answerability"
    label = f"{vector['id']}: {vector['description']}"
    assert check_answerability(vector["form"]) == vector["expectedViolations"], label


@pytest.mark.parametrize("path", VECTORS, ids=lambda p: p.stem)
def test_the_publish_gate_agrees_with_the_vector(path: pathlib.Path) -> None:
    """The check is only worth having if the publish path actually runs it."""
    vector = json.loads(path.read_text())
    expected = vector["expectedViolations"]

    if not expected:
        check_publishable(vector["form"])  # must not raise
        return

    with pytest.raises(PublishRefused) as refusal:
        check_publishable(vector["form"])
    assert refusal.value.violations == expected, vector["id"]


# --------------------------------------------------------------------------
# The importer's half: the same refusal, against a cell
# --------------------------------------------------------------------------
#
# A builder has no cell and gets the vector's sentence; an author with a
# spreadsheet open should not have to hunt for the row. Both come from
# `answerability.required_valueless_message`, one definition, for the reason
# defect 28 established: two texts is how two gates come to decide differently.


def _workbook(rows: list[list[str]]) -> bytes:
    import io

    import openpyxl

    book = openpyxl.Workbook()
    survey = book.active
    assert survey is not None
    survey.title = "survey"
    survey.append(["type", "name", "label::English (en)", "required"])
    for row in rows:
        survey.append(row + [""] * (4 - len(row)))
    choices = book.create_sheet("choices")
    choices.append(["list_name", "name", "label::English (en)"])
    settings = book.create_sheet("settings")
    settings.append(["form_title", "form_id", "version", "default_language"])
    settings.append(["Answerability", "answerability_test", "1", "English (en)"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_the_importer_refuses_a_required_note_against_its_row() -> None:
    from app.modules.forms.xlsform.importer import import_workbook

    result = import_workbook(
        _workbook(
            [
                ["text", "name", "Name", "yes"],
                ["note", "intro", "Read this out", "yes"],
            ]
        )
    )
    refusals = [d for d in result.diagnostics if d.code == "required_valueless_question"]
    assert len(refusals) == 1
    assert not result.publishable
    assert refusals[0].ref is not None
    assert refusals[0].ref.column == "required", "the cell, not just the form"
    assert refusals[0].message == check_answerability(result.form)[0], (
        "the importer and the publish gate say one sentence"
    )


def test_a_required_no_on_a_note_is_a_warning_and_not_a_refusal() -> None:
    """The severity follows the document, because the rule is about the document.

    `required: no` puts no `required` into the IR at all, so the document is
    clean and the publish gate would take it. An importer that refused it would
    be enforcing a stricter rule than §10.2 — the same importer-versus-gate
    disagreement defect 28 was, pointing the other way. It is still worth
    saying, because `no` is one keystroke from `yes`.
    """
    from app.modules.forms.xlsform.importer import import_workbook

    result = import_workbook(
        _workbook([["note", "intro", "Read this out", "no"], ["text", "name", "Name"]])
    )
    assert [d.code for d in result.diagnostics if d.severity == "error"] == []
    assert [d.code for d in result.diagnostics if d.severity == "warning"] == [
        "required_on_valueless_question"
    ]
    assert result.publishable
    assert check_answerability(result.form) == [], (
        "the importer and the gate agree about the document it produced"
    )



def test_a_note_with_no_required_column_imports_clean() -> None:
    from app.modules.forms.xlsform.importer import import_workbook

    result = import_workbook(
        _workbook([["note", "intro", "Read this out"], ["text", "name", "Name", "yes"]])
    )
    assert [d.message for d in result.diagnostics if d.severity == "error"] == []
    assert result.publishable
